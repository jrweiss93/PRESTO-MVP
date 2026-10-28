"""Genera la presentación de Presto en PowerPoint y HTML.

Uso:
    python scripts/build_presentation.py

Salidas:
    docs/presentacion-presto.pptx
    docs/presentacion-presto.html   (para imprimir el PDF con Chrome)
"""

from html import escape
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / 'docs'

NAVY = RGBColor(0x1E, 0x3A, 0x8A)
BRAND = RGBColor(0x46, 0x5F, 0xFF)
LIGHT = RGBColor(0xF2, 0xF7, 0xFF)
PALE = RGBColor(0xE0, 0xE7, 0xFF)
DARK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x6B, 0x72, 0x80)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
AMBER = RGBColor(0x92, 0x40, 0x0E)
AMBER_BG = RGBColor(0xFE, 0xF3, 0xC7)
SOFT_BLUE = RGBColor(0xDD, 0xE9, 0xFF)

DEMO_URL = 'https://presto-mvp-alpha.vercel.app'
REPO_URL = 'https://github.com/jrweiss93/PRESTO-MVP'
DEMO_USERS = 'ana@presto.demo  ·  beto@presto.demo  ·  carla@presto.demo  ·  Presto123!'


SLIDES = [
    {
        'kind': 'cover',
        'title': 'Presto',
        'subtitle': 'Finanzas e inversiones entre personas, en pesos y dólares.',
        'tag': 'DEMO PÚBLICA · DINERO SIMULADO',
        'url': DEMO_URL,
        'date': 'Octubre 2026',
        'notes': (
            'Presto es una aplicación que reúne billetera, préstamos entre personas '
            'e inversiones en un solo lugar. La versión que van a ver es una demo '
            'pública con dinero simulado: no se mueve plata real.'
        ),
    },
    {
        'kind': 'bullets',
        'title': 'El problema',
        'kicker': '01 · CONTEXTO',
        'bullets': [
            'Pedir un préstamo suele ser lento, con requisitos poco claros y costos que aparecen al final.',
            'Ahorrar en pesos es difícil y la inversión parece reservada a especialistas.',
            'Las herramientas están separadas: banco, billetera, planilla y app de inversión.',
            'Quien tiene dinero para prestar no encuentra un lugar simple con reglas claras.',
        ],
        'note': 'La claridad también es parte del producto.',
        'notes': (
            'El problema no es solo la plata: es la fricción y la falta de claridad. '
            'La gente no sabe cuánto va a pagar, dónde está su dinero ni cómo empezar a invertir.'
        ),
    },
    {
        'kind': 'cards',
        'title': 'La solución',
        'kicker': '02 · PROPUESTA',
        'intro': 'Una sola aplicación para mover, prestar, invertir y ordenar.',
        'cards': [
            ('Billetera en dos monedas', 'Saldo en ARS y USD, depósitos, retiros y transferencias.'),
            ('Préstamos entre personas', 'Publicá una oferta o elegí la que mejor se adapta a vos.'),
            ('Inversiones simples', 'Plazos fijos con tasa, interés y vencimiento claros.'),
            ('Orden financiero', 'Presupuestos por categoría y control de gastos.'),
        ],
        'cols': 2,
        'note': 'Presti: un asistente que explica el mercado, tus cuotas y el marco legal sin tecnicismos.',
        'notes': (
            'La propuesta es juntar lo que normalmente está disperso. '
            'El usuario no necesita cuatro aplicaciones ni una planilla paralela.'
        ),
    },
    {
        'kind': 'personas',
        'title': '¿A quién está dirigida?',
        'kicker': '03 · PÚBLICO',
        'cards': [
            ('Ana, 29', 'Prestataria\nNecesita financiarse para un proyecto y quiere saber exactamente cuánto va a pagar por mes.'),
            ('Beto, 45', 'Prestamista\nTiene saldo y quiere prestarlo con reglas claras, sin perder el control de su dinero.'),
            ('Carla, 35', 'Inversora\nQuiere que su saldo rinda sin volverse experta en mercados ni en jerga financiera.'),
        ],
        'note': 'Personas sin experiencia financiera que valoran claridad, orden y una interfaz simple.',
        'cols': 3,
        'notes': (
            'Son tres perfiles concretos: quien necesita dinero, quien lo presta y quien invierte. '
            'La app está pensada para gente común, no para traders.'
        ),
    },
    {
        'kind': 'cards',
        'title': '¿Qué operaciones permite?',
        'kicker': '04 · PRODUCTO',
        'cards': [
            ('Identidad', 'Verificación de identidad antes de operar con dinero.'),
            ('Billetera', 'Depositar, retirar y transferir en ARS o USD.'),
            ('Ofertas', 'Publicar, buscar y comparar ofertas de préstamo.'),
            ('Préstamos', 'Aceptar una oferta, recibir el dinero y pagar cuotas.'),
            ('Inversiones', 'Abrir un plazo fijo y cobrar el interés al vencimiento.'),
            ('Control', 'Presupuestos por categoría y seguimiento de gastos.'),
        ],
        'cols': 3,
        'note': 'Cada operación queda registrada y se puede seguir desde el panel.',
        'notes': (
            'El circuito completo: identidad, billetera, oferta, préstamo, cuotas e inversión. '
            'Cada operación deja registro y se puede seguir desde el panel.'
        ),
    },
    {
        'kind': 'flow',
        'title': 'El circuito completo',
        'kicker': '05 · CÓMO FUNCIONA',
        'steps': [
            ('1', 'Verificás tu identidad', 'KYC una vez y ya podés operar.'),
            ('2', 'Cargás saldo', 'Depósito simulado en la demo.'),
            ('3', 'Publicás o elegís', 'Oferta propia o una del marketplace.'),
            ('4', 'Pagás en cuotas', 'Cronograma claro, mes a mes.'),
            ('5', 'Invertís', 'Plazo fijo con vencimiento e interés.'),
        ],
        'note': 'El circuito completo se puede mostrar en cinco minutos.',
        'notes': (
            'Este es el recorrido que se puede mostrar en cinco minutos. '
            'La persona siempre sabe en qué paso está y qué sigue.'
        ),
    },
    {
        'kind': 'table',
        'title': 'Qué cambia con Presto',
        'kicker': '06 · BENEFICIOS',
        'headers': ['Lo habitual', 'En Presto'],
        'rows': [
            ('Una app para la cuenta, otra para invertir, una planilla para los gastos', 'Todo el circuito en una sola aplicación'),
            ('Costos que aparecen al final', 'Cuota, interés y costo total antes de aceptar'),
            ('Inversión con jerga y requisitos altos', 'Tasas, plazos y vencimientos explicados en simple'),
            ('Una sola moneda o conversión automática', 'ARS y USD separados, sin conversión implícita'),
            ('Sin asistencia para entender', 'Presti explica el mercado y el marco legal'),
        ],
        'note': 'Comparación de enfoque y experiencia de uso. No compara precios ni condiciones de competidores.',
        'notes': (
            'Las ventajas no son solo features: son menos pasos, menos sorpresas y más claridad. '
            'No decimos que los demás sean malos; decimos qué problema resolvemos distinto.'
        ),
    },
    {
        'kind': 'cards',
        'title': 'Presti, el asistente',
        'kicker': '07 · ASISTENTE',
        'cards': [
            ('Tu cuenta', 'Saldo, cuotas y próximos vencimientos.'),
            ('Mercado', 'Dólar y cripto con datos reales.'),
            ('Marco legal', 'Fichas simples con normas y enlaces.'),
            ('Sin inventos', 'Reglas y contenido curado; no IA generativa.'),
        ],
        'cols': 2,
        'note': 'Si no entiende una pregunta, propone opciones. No promete asesoramiento ni inventa datos.',
        'notes': (
            'Presti nació para explicar, no para impresionar. '
            'Responde con datos de la cuenta y textos legales revisados. No es un chat que alucina.'
        ),
    },
    {
        'kind': 'cards',
        'title': 'Confianza en cada movimiento',
        'kicker': '08 · SEGURIDAD',
        'cards': [
            ('Identidad verificada', 'KYC antes de operar con dinero.'),
            ('Cada movimiento, dos veces', 'Libro contable de doble entrada.'),
            ('Sin cobros duplicados', 'Cada operación tiene una clave única.'),
            ('Monedas separadas', 'ARS y USD nunca se mezclan.'),
        ],
        'cols': 2,
        'notes': (
            'La confianza se construye con reglas, no con promesas. '
            'Cada movimiento queda registrado y el sistema impide inconsistencias.'
        ),
    },
    {
        'kind': 'cards',
        'title': '¿Por qué estas herramientas?',
        'kicker': '09 · TECNOLOGÍA',
        'cards': [
            ('Python + Django', 'Orden, seguridad y velocidad para reglas de dinero.'),
            ('PostgreSQL', 'Consistencia y transacciones para datos financieros.'),
            ('Redis + Celery', 'Tareas diarias sin bloquear la aplicación.'),
            ('React + Vite', 'Interfaz rápida y pensada primero para celular.'),
            ('Docker + Render + Vercel', 'Despliegue reproducible y demo accesible.'),
        ],
        'cols': 3,
        'note': 'Se eligieron herramientas probadas y mantenibles, no las de moda.',
        'notes': (
            'La prioridad fue que el dinero esté bien cuidado y que el equipo pueda mantenerlo. '
            'Por eso Django y PostgreSQL, y una interfaz liviana con React.'
        ),
    },
    {
        'kind': 'stats',
        'title': 'Estado actual',
        'kicker': '10 · HOY',
        'stats': [
            ('13', 'módulos'),
            ('800', 'pruebas automáticas'),
            ('5', 'controles en cada cambio'),
            ('1', 'demo pública'),
        ],
        'bullets': [
            'Login, KYC, billetera, préstamos, inversiones, presupuestos, Presti y admin.',
            'Demo pública con dinero simulado y datos de ejemplo.',
            'Cada cambio pasa por calidad, migraciones, tests, front e imagen Docker.',
        ],
        'note': 'Demo pública: https://presto-mvp-alpha.vercel.app',
        'notes': (
            'No es un prototipo visual: es una aplicación completa con pruebas. '
            'La demo ya muestra el circuito de punta a punta.'
        ),
    },
    {
        'kind': 'flow',
        'title': 'Próximos pasos',
        'kicker': '11 · ROADMAP',
        'steps': [
            ('1', 'Pagos reales', 'Mercado Pago sandbox y luego producción.'),
            ('2', 'KYC con documentos', 'Subida y validación de identidad.'),
            ('3', 'Legal', 'TyC, privacidad y encuadre regulatorio.'),
            ('4', 'Operación', 'Sentry, backups y dominio propio.'),
        ],
        'note': 'La demo está lista; esto falta para usuarios reales.',
        'notes': (
            'La demo está lista. Para usuarios reales falta conectar pagos, validar documentos, '
            'cerrar el marco legal y sumar monitoreo.'
        ),
    },
    {
        'kind': 'closing',
        'title': 'Presto',
        'subtitle': 'Menos fricción, más claridad.',
        'url': DEMO_URL,
        'meta': f'Usuarios demo: {DEMO_USERS}  ·  Repo: {REPO_URL}',
        'notes': (
            'Cierre: Presto ya muestra el circuito completo. '
            'El siguiente paso es llevar la demo a una operación real, con pagos y legales.'
        ),
    },
]


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------


def add_text(
    slide,
    left,
    top,
    width,
    height,
    text,
    size=14,
    bold=False,
    color=DARK,
    align=PP_ALIGN.LEFT,
    anchor=MSO_ANCHOR.TOP,
    line_spacing=1.0,
):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    frame = box.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = anchor
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.line_spacing = line_spacing
    run = paragraph.add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = 'Segoe UI'
    run.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, fill, line=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    rect = slide.shapes.add_shape(shape, Inches(left), Inches(top), Inches(width), Inches(height))
    rect.fill.solid()
    rect.fill.fore_color.rgb = fill
    if line is None:
        rect.line.fill.background()
    else:
        rect.line.color.rgb = line
        rect.line.width = Pt(1)
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        rect.adjustments[0] = 0.08
    rect.text_frame.word_wrap = True
    return rect


def add_footer(slide, number):
    add_text(slide, 0.55, 7.1, 8, 0.22, 'Presto · demo pública · 2026', 9, False, GRAY)
    add_text(slide, 12.0, 7.1, 0.75, 0.22, str(number), 9, False, GRAY, PP_ALIGN.RIGHT)


def add_base(prs, title, number, kicker=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = WHITE
    add_text(slide, 0.55, 0.33, 3, 0.25, 'Presto.', 11, True, BRAND)
    if kicker:
        add_text(slide, 8.2, 0.33, 4.55, 0.25, kicker, 9, True, GRAY, PP_ALIGN.RIGHT)
    add_text(slide, 0.55, 0.72, 12.2, 0.72, title, 30, True, NAVY, line_spacing=1.05)
    add_rect(slide, 0.55, 1.52, 1.35, 0.06, BRAND, shape=MSO_SHAPE.RECTANGLE)
    add_footer(slide, number)
    return slide


def add_card(slide, left, top, width, height, title, body):
    card = add_rect(slide, left, top, width, height, LIGHT, line=PALE)
    frame = card.text_frame
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    frame.margin_top = Inches(0.13)
    frame.margin_bottom = Inches(0.08)
    paragraph = frame.paragraphs[0]
    paragraph.text = title
    paragraph.space_after = Pt(4)
    run = paragraph.runs[0]
    run.font.size = Pt(15)
    run.font.bold = True
    run.font.name = 'Segoe UI'
    run.font.color.rgb = NAVY
    for line in body.split('\n'):
        paragraph = frame.add_paragraph()
        paragraph.text = line
        paragraph.space_after = Pt(1)
        run = paragraph.runs[0]
        run.font.size = Pt(11.5)
        run.font.name = 'Segoe UI'
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)


def add_cards(slide, cards, cols, top=1.85, height=1.45, gap=0.22):
    width = (12.2 - gap * (cols - 1)) / cols
    for index, (title, body) in enumerate(cards):
        row, col = divmod(index, cols)
        add_card(
            slide,
            0.55 + col * (width + gap),
            top + row * (height + gap),
            width,
            height,
            title,
            body,
        )


def add_note(slide, text, top=5.55):
    shape = add_rect(slide, 0.55, top, 12.2, 0.52, AMBER_BG)
    frame = shape.text_frame
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    frame.margin_left = Inches(0.16)
    frame.margin_right = Inches(0.16)
    paragraph = frame.paragraphs[0]
    paragraph.text = text
    run = paragraph.runs[0]
    run.font.size = Pt(11.5)
    run.font.bold = True
    run.font.name = 'Segoe UI'
    run.font.color.rgb = AMBER


def add_table(slide, headers, rows, top=1.9):
    shape = slide.shapes.add_table(
        len(rows) + 1, len(headers), Inches(0.55), Inches(top), Inches(12.2), Inches(0.5 * (len(rows) + 1))
    )
    table = shape.table
    table.columns[0].width = Inches(6.2)
    table.columns[1].width = Inches(6.0)
    table.rows[0].height = Inches(0.45)
    for row in table.rows:
        row.height = Inches(0.6)
    for col, header in enumerate(headers):
        cell = table.cell(0, col)
        cell.text = header
        cell.fill.solid()
        cell.fill.fore_color.rgb = BRAND
        cell.margin_left = Inches(0.08)
        cell.margin_right = Inches(0.08)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(11.5)
                run.font.bold = True
                run.font.name = 'Segoe UI'
                run.font.color.rgb = WHITE
    for row_index, row in enumerate(rows, start=1):
        for col_index, value in enumerate(row):
            cell = table.cell(row_index, col_index)
            cell.text = value
            cell.fill.solid()
            cell.fill.fore_color.rgb = WHITE if row_index % 2 else RGBColor(0xF8, 0xFA, 0xFC)
            cell.margin_left = Inches(0.08)
            cell.margin_right = Inches(0.08)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(10.5)
                    run.font.name = 'Segoe UI'
                    run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)


def add_flow(slide, steps, top=2.3):
    count = len(steps)
    gap = 0.28
    width = (12.2 - gap * (count - 1)) / count
    for index, (number, title, body) in enumerate(steps):
        left = 0.55 + index * (width + gap)
        card = add_rect(slide, left, top, width, 2.25, LIGHT, line=PALE)
        frame = card.text_frame
        frame.vertical_anchor = MSO_ANCHOR.TOP
        frame.margin_left = Inches(0.14)
        frame.margin_right = Inches(0.14)
        frame.margin_top = Inches(0.12)
        paragraph = frame.paragraphs[0]
        paragraph.text = number
        run = paragraph.runs[0]
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.color.rgb = BRAND
        run.font.name = 'Segoe UI'
        paragraph = frame.add_paragraph()
        paragraph.text = title
        paragraph.space_before = Pt(2)
        run = paragraph.runs[0]
        run.font.size = Pt(12.5)
        run.font.bold = True
        run.font.color.rgb = NAVY
        run.font.name = 'Segoe UI'
        paragraph = frame.add_paragraph()
        paragraph.text = body
        run = paragraph.runs[0]
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
        run.font.name = 'Segoe UI'
        if index < count - 1:
            add_text(
                slide,
                left + width + 0.02,
                top + 0.72,
                0.24,
                0.35,
                '→',
                18,
                True,
                BRAND,
                PP_ALIGN.CENTER,
            )


def add_wordmark(slide, left, top, size=58):
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(8.5), Inches(1.05))
    frame = box.text_frame
    frame.word_wrap = True
    frame.margin_left = 0
    frame.margin_right = 0
    frame.margin_top = 0
    frame.margin_bottom = 0
    paragraph = frame.paragraphs[0]
    paragraph.alignment = PP_ALIGN.LEFT
    run = paragraph.add_run()
    run.text = 'Presto'
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.name = 'Segoe UI'
    run.font.color.rgb = WHITE
    dot = paragraph.add_run()
    dot.text = '.'
    dot.font.size = Pt(size)
    dot.font.bold = True
    dot.font.name = 'Segoe UI'
    dot.font.color.rgb = RGBColor(0x93, 0xC5, 0xFD)


def add_cover(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    mark = add_rect(slide, 0.85, 1.0, 0.95, 0.95, BRAND)
    mark.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = mark.text_frame.paragraphs[0]
    paragraph.text = 'P'
    paragraph.alignment = PP_ALIGN.CENTER
    run = paragraph.runs[0]
    run.font.size = Pt(40)
    run.font.bold = True
    run.font.color.rgb = WHITE
    run.font.name = 'Segoe UI'
    add_text(slide, 2.1, 2.0, 9.6, 0.8, data['subtitle'], 20, False, SOFT_BLUE, line_spacing=1.2)
    add_wordmark(slide, 2.05, 1.02, 58)
    tag = add_rect(slide, 0.85, 3.35, 4.65, 0.42, AMBER_BG)
    tag.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = tag.text_frame.paragraphs[0]
    paragraph.text = data['tag']
    run = paragraph.runs[0]
    run.font.size = Pt(10.5)
    run.font.bold = True
    run.font.color.rgb = AMBER
    run.font.name = 'Segoe UI'
    add_text(slide, 0.85, 4.0, 10.8, 0.5, 'Una demo para mostrar el circuito completo, sin dinero real.', 15, False, SOFT_BLUE)
    add_text(slide, 0.85, 6.15, 10, 0.3, data['url'], 13, True, WHITE)
    add_text(slide, 0.85, 6.5, 10, 0.3, data['date'], 10, False, SOFT_BLUE)
    slide.notes_slide.notes_text_frame.text = data['notes']


def add_closing(prs, data):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    add_wordmark(slide, 0.85, 1.2, 56)
    add_text(slide, 0.85, 2.15, 11.8, 0.7, data['subtitle'], 26, False, SOFT_BLUE)
    add_text(slide, 0.85, 3.35, 11.8, 0.35, data['url'], 15, True, WHITE)
    add_text(slide, 0.85, 3.8, 11.8, 0.6, data['meta'], 10.5, False, SOFT_BLUE, line_spacing=1.3)
    add_text(slide, 0.85, 6.4, 11.8, 0.3, 'Gracias.', 16, True, WHITE)
    slide.notes_slide.notes_text_frame.text = data['notes']


def build_pptx(slides=None, output=None):
    slides = slides or SLIDES
    output = Path(output) if output else DOCS / 'presentacion-presto.pptx'
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for index, data in enumerate(slides, start=1):
        kind = data['kind']
        if kind == 'cover':
            add_cover(prs, data)
            continue
        if kind == 'closing':
            add_closing(prs, data)
            continue

        slide = add_base(prs, data['title'], index, data.get('kicker'))
        if kind == 'bullets':
            add_text(slide, 0.7, 1.9, 11.9, 4.8, '\n'.join(f'•  {item}' for item in data['bullets']), 17, False, DARK, line_spacing=1.35)
            if data.get('note'):
                add_note(slide, data['note'], top=5.95)
        elif kind in {'cards', 'personas'}:
            if data.get('intro'):
                add_text(slide, 0.55, 1.75, 12.0, 0.4, data['intro'], 15, False, GRAY)
            top = 2.2 if data.get('intro') else 2.0
            height = 2.9 if kind == 'personas' else (1.7 if len(data['cards']) > 4 else 1.6)
            add_cards(slide, data['cards'], data.get('cols', 2), top=top, height=height)
            if data.get('note'):
                add_note(slide, data['note'], top=5.5 if kind == 'personas' else 5.95)
        elif kind == 'flow':
            add_flow(slide, data['steps'])
            if data.get('note'):
                add_note(slide, data['note'], top=5.7)
        elif kind == 'table':
            add_table(slide, data['headers'], data['rows'])
            add_note(slide, data['note'], top=5.95)
        elif kind == 'stats':
            stats = data['stats']
            width = (12.2 - 0.22 * 3) / 4
            for index_stat, (number, label) in enumerate(stats):
                left = 0.55 + index_stat * (width + 0.22)
                card = add_rect(slide, left, 2.0, width, 1.75, LIGHT, line=PALE)
                frame = card.text_frame
                frame.vertical_anchor = MSO_ANCHOR.MIDDLE
                paragraph = frame.paragraphs[0]
                paragraph.text = number
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.runs[0]
                run.font.size = Pt(34)
                run.font.bold = True
                run.font.color.rgb = BRAND
                run.font.name = 'Segoe UI'
                paragraph = frame.add_paragraph()
                paragraph.text = label
                paragraph.alignment = PP_ALIGN.CENTER
                run = paragraph.runs[0]
                run.font.size = Pt(10.5)
                run.font.color.rgb = GRAY
                run.font.name = 'Segoe UI'
            add_text(slide, 0.7, 4.0, 11.9, 2.5, '\n'.join(f'•  {item}' for item in data['bullets']), 15, False, DARK, line_spacing=1.35)
            if data.get('note'):
                add_note(slide, data['note'], top=6.3)

        slide.notes_slide.notes_text_frame.text = data['notes']

    output = Path(output) if output else DOCS / 'presentacion-presto.pptx'
    prs.save(output)
    return output


# ---------------------------------------------------------------------------
# HTML para exportar a PDF
# ---------------------------------------------------------------------------


def html_card(title, body):
    body_html = '<br>'.join(escape(line) for line in body.split('\n'))
    return f'<div class="card"><h3>{escape(title)}</h3><p>{body_html}</p></div>'


def html_slide(data, number):
    kind = data['kind']
    footer = f'<div class="footer"><span>Presto · demo pública · 2026</span><span>{number}</span></div>'

    if kind == 'cover':
        return f'''
        <section class="slide cover">
          <div class="pmark">P</div>
          <h1>Presto<span class="dot">.</span></h1>
          <div class="sub">{escape(data['subtitle'])}</div>
          <div class="badge">{escape(data['tag'])}</div>
          <p class="cover-text">Una demo para mostrar el circuito completo, sin dinero real.</p>
          <div class="meta"><div class="url">{escape(data['url'])}</div><div>{escape(data['date'])}</div></div>
        </section>'''

    if kind == 'closing':
        return f'''
        <section class="slide closing">
          <h1>Presto<span class="dot">.</span></h1>
          <div class="sub">{escape(data['subtitle'])}</div>
          <div class="url">{escape(data['url'])}</div>
          <div class="meta">{escape(data['meta'])}</div>
          <div class="thanks">Gracias.</div>
        </section>'''

    kicker = f'<div class="kicker">{escape(data["kicker"])}</div>' if data.get('kicker') else ''
    content = ''
    if kind == 'bullets':
        content = '<ul class="bullets">' + ''.join(f'<li>{escape(item)}</li>' for item in data['bullets']) + '</ul>'
        if data.get('note'):
            content += f'<div class="note">{escape(data["note"])}</div>'
    elif kind in {'cards', 'personas'}:
        intro = f'<p class="intro">{escape(data["intro"])}</p>' if data.get('intro') else ''
        cols = data.get('cols', 2)
        cards = ''.join(html_card(title, body) for title, body in data['cards'])
        note = f'<div class="note">{escape(data["note"])}</div>' if data.get('note') else ''
        content = f'{intro}<div class="cards cols{cols}">{cards}</div>{note}'
    elif kind == 'flow':
        steps = ''.join(
            f'<div class="step"><div class="n">{escape(number)}</div><h3>{escape(title)}</h3><p>{escape(body)}</p></div>'
            for number, title, body in data['steps']
        )
        note = f'<div class="note">{escape(data["note"])}</div>' if data.get('note') else ''
        content = f'<div class="flow">{steps}</div>{note}'
    elif kind == 'table':
        head = ''.join(f'<th>{escape(header)}</th>' for header in data['headers'])
        rows = ''.join(
            '<tr>' + ''.join(f'<td>{escape(value)}</td>' for value in row) + '</tr>'
            for row in data['rows']
        )
        content = f'<table><thead><tr>{head}</tr></thead><tbody>{rows}</tbody></table><div class="note">{escape(data["note"])}</div>'
    elif kind == 'stats':
        stats = ''.join(
            f'<div class="stat"><div class="num">{escape(number)}</div><div class="lbl">{escape(label)}</div></div>'
            for number, label in data['stats']
        )
        bullets = '<ul class="bullets small">' + ''.join(f'<li>{escape(item)}</li>' for item in data['bullets']) + '</ul>'
        note = f'<div class="note">{escape(data["note"])}</div>' if data.get('note') else ''
        content = f'<div class="stats">{stats}</div>{bullets}{note}'

    return f'''
    <section class="slide">
      <div class="brand">Presto.</div>{kicker}
      <h1 class="title">{escape(data['title'])}</h1>
      <div class="rule"></div>
      <div class="content">{content}</div>
      {footer}
    </section>'''


HTML_CSS = '''
:root {
  --navy: #1e3a8a;
  --brand: #465fff;
  --light: #f2f7ff;
  --pale: #e0e7ff;
  --dark: #111827;
  --gray: #6b7280;
  --amber: #92400e;
  --amber-bg: #fef3c7;
}
@page { size: 13.333in 7.5in; margin: 0; }
* { box-sizing: border-box; }
body { margin: 0; font-family: "Segoe UI", Calibri, Arial, sans-serif; color: var(--dark); }
.slide {
  position: relative;
  width: 13.333in;
  height: 7.5in;
  padding: 0.45in 0.6in 0.5in;
  background: #fff;
  page-break-after: always;
  overflow: hidden;
  -webkit-print-color-adjust: exact;
  print-color-adjust: exact;
}
.slide:last-child { page-break-after: auto; }
.brand { color: var(--brand); font-weight: 700; font-size: 11pt; }
.kicker { position: absolute; top: 0.45in; right: 0.6in; font-size: 9pt; color: var(--gray); letter-spacing: 0.08em; font-weight: 600; }
h1.title { font-size: 30pt; line-height: 1.1; color: var(--navy); margin: 0.22in 0 0.12in; font-weight: 700; }
.rule { width: 1.3in; height: 0.06in; background: var(--brand); border-radius: 3px; }
.content { margin-top: 0.28in; }
.intro { margin: 0 0 0.16in; font-size: 15pt; color: var(--gray); }
ul.bullets { margin: 0; padding-left: 0.32in; font-size: 17pt; line-height: 1.45; }
ul.bullets li { margin: 0.1in 0; }
ul.bullets.small { font-size: 14pt; margin-top: 0.2in; }
.cards { display: grid; gap: 0.2in; }
.cards.cols2 { grid-template-columns: 1fr 1fr; }
.cards.cols3 { grid-template-columns: 1fr 1fr 1fr; }
.card { background: var(--light); border: 1px solid var(--pale); border-radius: 0.12in; padding: 0.16in 0.18in; }
.card h3 { margin: 0 0 0.06in; font-size: 14.5pt; color: var(--navy); }
.card p { margin: 0; font-size: 11.5pt; line-height: 1.35; color: #4b5563; }
.note { margin-top: 0.2in; background: var(--amber-bg); color: var(--amber); border-radius: 0.1in; padding: 0.12in 0.16in; font-size: 12pt; font-weight: 600; }
table { width: 100%; border-collapse: collapse; font-size: 12pt; margin: 0.05in 0 0.15in; }
th { background: var(--brand); color: #fff; text-align: left; padding: 0.1in; font-size: 11.5pt; }
td { border-bottom: 1px solid #e5e7eb; padding: 0.1in; vertical-align: top; color: #374151; }
tr:nth-child(even) td { background: #f8fafc; }
.stats { display: grid; grid-template-columns: repeat(4, 1fr); gap: 0.2in; margin-bottom: 0.2in; }
.stat { background: var(--light); border: 1px solid var(--pale); border-radius: 0.12in; padding: 0.18in; text-align: center; }
.stat .num { font-size: 34pt; font-weight: 700; color: var(--brand); line-height: 1; }
.stat .lbl { font-size: 10.5pt; color: var(--gray); margin-top: 0.08in; }
.flow { display: flex; align-items: stretch; gap: 0.12in; margin-top: 0.35in; }
.flow .step { flex: 1; background: var(--light); border: 1px solid var(--pale); border-radius: 0.12in; padding: 0.14in; }
.flow .step .n { font-size: 10pt; font-weight: 700; color: var(--brand); }
.flow .step h3 { margin: 0.04in 0; font-size: 13pt; color: var(--navy); }
.flow .step p { margin: 0; font-size: 10.5pt; color: #4b5563; line-height: 1.3; }
.footer { position: absolute; bottom: 0.25in; left: 0.6in; right: 0.6in; display: flex; justify-content: space-between; font-size: 9pt; color: var(--gray); }
.slide.cover, .slide.closing { background: var(--navy); color: #fff; }
.cover .pmark { width: 0.9in; height: 0.9in; border-radius: 0.18in; background: var(--brand); color: #fff; font-size: 38pt; font-weight: 700; display: flex; align-items: center; justify-content: center; margin-bottom: 0.22in; }
.cover h1 { font-size: 62pt; margin: 0 0 0.05in; color: #fff; }
.cover .dot, .closing .dot { color: #93c5fd; }
.cover .sub { font-size: 22pt; color: #dbeafe; max-width: 9.2in; line-height: 1.3; }
.cover .badge { display: inline-block; background: var(--amber-bg); color: var(--amber); font-size: 10.5pt; font-weight: 700; letter-spacing: 0.06em; padding: 0.08in 0.14in; border-radius: 0.08in; margin-top: 0.35in; }
.cover-text { color: #c7d2fe; font-size: 15pt; margin-top: 0.2in; }
.cover .meta { position: absolute; bottom: 0.55in; left: 0.6in; font-size: 11pt; color: #c7d2fe; }
.cover .url { color: #fff; font-size: 14pt; font-weight: 600; margin-bottom: 0.06in; }
.closing h1 { font-size: 56pt; color: #fff; margin: 0.1in 0; }
.closing .sub { font-size: 26pt; color: #dbeafe; }
.closing .url { color: #fff; font-size: 15pt; margin-top: 0.45in; }
.closing .meta { color: #c7d2fe; font-size: 10.5pt; margin-top: 0.15in; line-height: 1.4; }
.closing .thanks { color: #fff; font-size: 16pt; font-weight: 600; margin-top: 0.55in; }
'''


def build_html(slides=None, output=None):
    slides = slides or SLIDES
    output = Path(output) if output else DOCS / 'presentacion-presto.html'
    deck = ''.join(html_slide(data, index) for index, data in enumerate(slides, start=1))
    document = f'''<!doctype html>
<html lang="es">
<head>
  <meta charset="utf-8">
  <title>Presto — presentación</title>
  <style>{HTML_CSS}</style>
</head>
<body>{deck}</body>
</html>'''
    output.write_text(document, encoding='utf-8')
    return output


if __name__ == '__main__':
    pptx_path = build_pptx()
    html_path = build_html()
    print(f'PPTX: {pptx_path}')
    print(f'HTML: {html_path}')
